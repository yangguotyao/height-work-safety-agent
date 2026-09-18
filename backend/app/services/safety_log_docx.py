from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

def _font(run, *, name: str = "宋体", size: float = 10.5, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def _text(paragraph, text: str, *, bold: bool = False, size: float = 10.5) -> None:
    _font(paragraph.add_run(str(text)), bold=bold, size=size)


def _format_cell(cell, *, header: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", 100), ("start", 120), ("bottom", 100), ("end", 120)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.08
        if header:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    size = 15 if level == 1 else 12
    _font(paragraph.add_run(text), name="黑体", size=size, bold=True)


def _bullet_list(document: Document, items: list[str], empty_text: str = "无") -> None:
    values = [str(item).strip() for item in items if str(item).strip()]
    if not values:
        paragraph = document.add_paragraph()
        _text(paragraph, empty_text)
        return
    for item in values:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        _text(paragraph, item)


def _table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = widths is None
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        _text(cell.paragraphs[0], header, bold=True, size=9.5)
        _format_cell(cell, header=True)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = ""
            _text(cells[index].paragraphs[0], value if value not in (None, "") else "—", size=9)
            _format_cell(cells[index])
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Cm(width)
            for row in table.rows:
                row.cells[index].width = Cm(width)


def _add_image(
    cell,
    image_path: Path,
    caption_text: str,
    *,
    width: float = 2.75,
) -> None:
    cell.text = ""
    if image_path.is_file():
        picture = cell.paragraphs[0]
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(image_path), width=Cm(width))
        caption = cell.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _text(caption, caption_text or image_path.name, size=7.5)
    else:
        _text(cell.paragraphs[0], caption_text or "暂无图片", size=8)
    _format_cell(cell)


def _basis_text(items: list[dict[str, Any]]) -> str:
    if not items:
        return "未关联规范条款"
    item = items[0]
    standard = item.get("standard_code") or item.get("standard_name") or "相关规范"
    name = item.get("standard_name")
    if name and name != standard:
        standard = f"{standard}《{name}》"
    clause = str(item.get("clause") or "").strip()
    quote = re.sub(r"\s+", "", str(item.get("quote") or "").strip())
    if len(quote) > 80:
        sentence_end = quote.find("。", 0, 81)
        quote = quote[: sentence_end + 1] if sentence_end >= 0 else f"{quote[:80]}……"
    return "\n".join(value for value in [f"{standard} {clause}".strip(), quote] if value)


def _onsite_closure_table(
    document: Document,
    inspections: list[dict[str, Any]],
    rectifications: list[dict[str, Any]],
    before_image_root: Path,
    after_image_root: Path,
) -> None:
    headers = ["整改前照片", "整改后照片", "隐患问题与规范依据", "整改与复核结果"]
    widths = [3.25, 3.25, 6.25, 3.85]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        _text(cell.paragraphs[0], header, bold=True, size=9.5)
        _format_cell(cell, header=True)
    inspections_by_id = {item.get("id"): item for item in inspections}
    linked_inspection_ids: set[str] = set()
    rows: list[tuple[dict[str, Any], dict[str, Any]]] = []
    for item in rectifications:
        inspection = inspections_by_id.get(item.get("inspection_id"), {})
        if item.get("inspection_id"):
            linked_inspection_ids.add(item["inspection_id"])
        rows.append((inspection, item))
    for inspection in inspections:
        if inspection.get("id") not in linked_inspection_ids:
            rows.append((inspection, {}))

    for inspection, item in rows:
        cells = table.add_row().cells
        before_name = item.get("before_stored_image_name") or inspection.get(
            "stored_image_name"
        )
        _add_image(
            cells[0],
            before_image_root / str(before_name or ""),
            item.get("before_original_filename")
            or inspection.get("original_filename")
            or "整改前照片",
        )
        _add_image(
            cells[1],
            after_image_root / str(item.get("after_stored_image_name") or ""),
            item.get("after_original_filename") or "尚未提交整改后照片",
        )

        issue_cell = cells[2]
        issue_cell.text = ""
        _text(
            issue_cell.paragraphs[0],
            " · ".join(filter(None, [item.get("item_no"), item.get("title")]))
            or "未形成正式安全事项",
            bold=True,
            size=9.5,
        )
        if item.get("fact_description"):
            paragraph = issue_cell.add_paragraph()
            _text(paragraph, str(item["fact_description"]), size=9)
        if item.get("rectification_requirement"):
            paragraph = issue_cell.add_paragraph()
            _text(paragraph, "整改要求：", bold=True, size=8.5)
            _text(paragraph, str(item["rectification_requirement"]), size=8.5)
        paragraph = issue_cell.add_paragraph()
        _text(paragraph, "相关规范依据：", bold=True, size=8.5)
        _text(paragraph, _basis_text(item.get("basis") or []), size=8.5)
        _format_cell(issue_cell)

        result_cell = cells[3]
        result_cell.text = ""
        status = {
            "closed": "已关闭",
            "pending_rectification": "待整改",
            "rectifying": "整改中",
            "pending_review": "待复核",
        }.get(item.get("order_status") or item.get("status"), item.get("order_status") or item.get("status"))
        _text(result_cell.paragraphs[0], status or "待人工确认", bold=True, size=9.5)
        if item.get("rectification_description"):
            paragraph = result_cell.add_paragraph()
            _text(paragraph, str(item["rectification_description"]), size=8.5)
        if item.get("latest_review_result"):
            paragraph = result_cell.add_paragraph()
            review = "复核通过" if item["latest_review_result"] == "pass" else "复核退回"
            _text(paragraph, review, bold=True, size=8.5)
            if item.get("latest_review_reason"):
                _text(paragraph, f"：{item['latest_review_reason']}", size=8.5)
        _format_cell(result_cell)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
        for row in table.rows:
            row.cells[index].width = Cm(width)


def render_safety_log_docx(log: dict[str, Any], target: Path) -> Path:
    """Render a restrained, black-and-white A4 report without LibreOffice."""
    target.parent.mkdir(parents=True, exist_ok=True)
    content = log["content"]
    summary = content["summary"]
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(28)
    title.paragraph_format.space_after = Pt(18)
    _font(title.add_run("施工安全全过程日志"), name="黑体", size=22, bold=True)

    project = document.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(project.add_run(log["project_name"]), name="黑体", size=14, bold=True)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _text(
        meta,
        f"日期：{log['assessment_date']}    版本：V{log['version']}    "
        f"生成时间：{log['created_at'][:19]}",
    )

    document.add_paragraph()
    _heading(document, "一、当日概况")
    _table(
        document,
        ["方案审查/修订", "班前风险分析", "现场巡检", "新增安全事项", "当日关闭", "未闭环事项"],
        [[
            summary.get("plan_activity_count", 0),
            summary["task_count"],
            summary.get("inspection_count", 0),
            summary.get("safety_item_count", 0),
            summary.get("closed_safety_item_count", 0),
            summary.get("open_item_count", 0),
        ]],
    )

    _heading(document, "二、当日安全过程时间线")
    timeline = content.get("timeline") or []
    if timeline:
        _table(
            document,
            ["时间", "阶段", "事件", "结果与状态"],
            [
                [
                    str(item.get("at") or "")[:19].replace("T", " "),
                    item.get("stage"),
                    item.get("title"),
                    "\n".join(filter(None, [item.get("detail"), item.get("status")])),
                ]
                for item in timeline
            ],
            widths=[3.1, 2.2, 4.1, 7.2],
        )
    else:
        _bullet_list(document, [], "当日暂无安全过程事件。")

    _heading(document, "三、施工方案审查与整改")
    plan_rows = []
    for item in content.get("plan_audits") or []:
        plan_rows.append(
            [
                item.get("filename"),
                "方案审查",
                f"发现{item.get('finding_count', 0)}项问题",
                "审查完成",
            ]
        )
    for item in content.get("plan_revisions") or []:
        comparison = item.get("comparison") or {}
        total = int(comparison.get("original_finding_count") or 0)
        resolved = int(comparison.get("resolved_count") or 0)
        remaining = max(total - resolved, 0)
        plan_rows.append(
            [
                item.get("revised_filename"),
                f"第{item.get('attempt_no')}次整改对比",
                f"原问题{total}项，已解决{resolved}项，仍需修改{remaining}项",
                "已关闭" if item.get("status") == "closed" else "需继续修订",
            ]
        )
    if plan_rows:
        _table(document, ["方案", "业务动作", "结果", "状态"], plan_rows)
    else:
        _bullet_list(document, [], "当日无施工方案审查或修订记录。")

    _heading(document, "四、班前风险分析")
    weather = content.get("weather") or {}
    if weather:
        paragraph = document.add_paragraph()
        _text(paragraph, "天气与环境：", bold=True)
        _text(paragraph, weather.get("summary") or "未获得有效天气摘要。")
    tasks = content.get("tasks") or []
    if not tasks:
        paragraph = document.add_paragraph()
        _text(paragraph, "当日暂无班前风险分析任务。")
    for index, task in enumerate(tasks, 1):
        _heading(document, f"4.{index} {task['normalized_task']}", level=2)
        _table(
            document,
            ["作业内容", "位置", "楼层/高度", "时间", "风险等级", "优先分"],
            [[
                task["work_content"],
                task["work_location"],
                task["work_floor"],
                task["work_time"],
                task["risk_level"],
                task["priority_score"],
            ]],
        )
        paragraph = document.add_paragraph()
        _text(paragraph, "风险说明：", bold=True)
        _text(paragraph, task.get("risk_summary") or "无")
        _heading(document, "主要风险", level=2)
        _bullet_list(document, task.get("main_risks") or [])
        _heading(document, "作业前检查要点", level=2)
        _bullet_list(document, task.get("pre_job_checks") or [])
        _heading(document, "禁止行为", level=2)
        _bullet_list(document, task.get("prohibited_behaviors") or [])
        _heading(document, "风险控制建议", level=2)
        _bullet_list(document, task.get("interventions") or [])
        if task.get("weather_warnings"):
            _heading(document, "天气与环境提示", level=2)
            _bullet_list(document, task.get("weather_warnings") or [])

    _heading(document, "五、现场隐患巡检与整改")
    inspections = content.get("onsite_inspections") or []
    rectifications = content.get("onsite_rectification") or []
    if inspections or rectifications:
        _onsite_closure_table(
            document,
            inspections,
            rectifications,
            target.parent.parent / "hazard_images",
            target.parent.parent / "rectification_images",
        )
    else:
        _bullet_list(document, [], "当日暂无现场多模态巡检记录。")

    _heading(document, "六、截至当日未闭环事项")
    open_items = content.get("open_items") or []
    _table(
        document,
        ["类型", "事项", "当前状态", "说明"],
        [
            [
                "方案整改" if item.get("type") == "plan_revision" else "现场隐患",
                item.get("title"),
                item.get("status"),
                item.get("detail"),
            ]
            for item in open_items
        ],
    ) if open_items else _bullet_list(document, [], "截至当日没有未闭环事项。")

    change = log.get("change") or {}
    if change.get("previous_log_id") and change.get("explanation"):
        paragraph = document.add_paragraph()
        _text(paragraph, change["explanation"])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(footer.add_run(f"{log['project_name']} · 施工安全全过程日志 V{log['version']}"), size=8)
    document.save(target)
    return target
