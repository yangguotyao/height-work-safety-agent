from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook

from ..db import Database, json_load
from ..repositories import Repository
from .retrieval import similarity

GOLD_HEADERS = [
    "高处场景",
    "方案内容",
    "来源位置",
    "问题描述",
    "审查依据",
    "证据来源",
    "风险后果",
    "标准答案（结果）",
    "修改建议",
]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def import_gold_cases(database: Database, workbook_path: Path) -> tuple[int, int]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"金标准工作簿不存在：{workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "审计金标准" not in workbook.sheetnames:
        raise ValueError("金标准工作簿缺少“审计金标准”工作表")
    worksheet = workbook["审计金标准"]
    header_row = None
    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if [_text(value) for value in row[:9]] == GOLD_HEADERS:
            header_row = row_index
            break
    if header_row is None:
        raise ValueError("金标准工作簿字段与首版数据契约不一致")

    now = datetime.now(timezone.utc).isoformat()
    rows = []
    skipped = 0
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        values = [_text(value) for value in row[:9]]
        if not values[0] or not values[7]:
            skipped += 1
            continue
        stable_id = uuid4().hex
        rows.append((stable_id, str(workbook_path), row_index, *values, now))

    database.executemany(
        """INSERT INTO gold_cases
           (id, source_workbook, source_row, scene, plan_quote, source_location, issue,
            basis, evidence_source, risk_consequence, expected_result, suggestion, imported_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(source_workbook, source_row) DO UPDATE SET
             scene=excluded.scene,
             plan_quote=excluded.plan_quote,
             source_location=excluded.source_location,
             issue=excluded.issue,
             basis=excluded.basis,
             evidence_source=excluded.evidence_source,
             risk_consequence=excluded.risk_consequence,
             expected_result=excluded.expected_result,
             suggestion=excluded.suggestion,
             imported_at=excluded.imported_at""",
        rows,
    )
    workbook.close()
    return len(rows), skipped


def _scene_score(gold_scene: str, item_scene: str) -> float:
    if gold_scene == item_scene:
        return 1.0
    if "脚手架" in gold_scene and "脚手架" in item_scene:
        return 0.8
    if any(term in gold_scene and term in item_scene for term in ("临边", "洞口", "栏杆", "平台")):
        return 0.7
    return similarity(gold_scene, item_scene)


class GoldEvaluator:
    def __init__(self, database: Database, repository: Repository):
        self.database = database
        self.repository = repository

    def resolve_source(self, audit_run_id: str, fallback_source: str) -> str:
        """Select the gold workbook that belongs to the uploaded document.

        Gold workbooks follow ``《document stem》...xlsx``.  Falling back keeps
        compatibility for ad-hoc documents that do not have a dedicated gold set.
        """
        run = self.repository.get_audit_run(audit_run_id)
        document = self.repository.get_document(run["document_id"])
        expected_prefix = f"《{Path(document['filename']).stem}》"
        sources = self.database.fetch_all(
            "SELECT DISTINCT source_workbook FROM gold_cases ORDER BY source_workbook"
        )
        matches = [
            row["source_workbook"]
            for row in sources
            if Path(row["source_workbook"]).stem.startswith(expected_prefix)
        ]
        if matches:
            return min(matches, key=lambda value: (len(Path(value).stem), value))
        return fallback_source

    def evaluate(self, audit_run_id: str, gold_source: str) -> dict[str, Any]:
        gold_source = self.resolve_source(audit_run_id, gold_source)
        gold_cases = self.database.fetch_all(
            "SELECT * FROM gold_cases WHERE source_workbook = ? ORDER BY source_row",
            (gold_source,),
        )
        if not gold_cases:
            raise ValueError("没有可用的金标准案例，请先导入金标准")
        items = self.repository.list_audit_items(audit_run_id)
        details = []
        matched_scores = []
        correct_results = 0

        for gold in gold_cases:
            best_item = None
            best_score = 0.0
            gold_text = f"{gold['plan_quote']} {gold['issue']}"
            for item in items:
                content_score = similarity(gold_text, f"{item['plan_quote']} {item['issue']}")
                scene_score = _scene_score(gold["scene"], item["scene"])
                score = 0.75 * content_score + 0.25 * scene_score
                if score > best_score:
                    best_score = score
                    best_item = item
            matched = best_item is not None and best_score >= 0.1
            actual_result = best_item["result"] if matched else None
            result_correct = bool(matched and actual_result == gold["expected_result"])
            if matched:
                matched_scores.append(best_score)
            if result_correct:
                correct_results += 1
            details.append(
                {
                    "gold_case_id": gold["id"],
                    "gold_scene": gold["scene"],
                    "expected_result": gold["expected_result"],
                    "audit_item_id": best_item["id"] if matched else None,
                    "actual_result": actual_result,
                    "match_score": round(best_score if matched else 0.0, 4),
                    "result_correct": result_correct,
                }
            )

        test_id = uuid4().hex
        now = datetime.now(timezone.utc).isoformat()
        total = len(gold_cases)
        result = {
            "id": test_id,
            "audit_run_id": audit_run_id,
            "gold_source": gold_source,
            "total_gold": total,
            "generated_items": len(items),
            "matched_cases": len(matched_scores),
            "correct_results": correct_results,
            "result_accuracy": round(correct_results / total, 4) if total else 0.0,
            "average_match_score": (
                round(sum(matched_scores) / len(matched_scores), 4) if matched_scores else 0.0
            ),
            "details": details,
            "created_at": now,
        }
        self.database.execute(
            """INSERT INTO test_runs
               (id, audit_run_id, gold_source, total_gold, generated_items, matched_cases,
                correct_results, result_accuracy, average_match_score, details_json, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                test_id,
                audit_run_id,
                gold_source,
                total,
                len(items),
                len(matched_scores),
                correct_results,
                result["result_accuracy"],
                result["average_match_score"],
                json.dumps(details, ensure_ascii=False),
                now,
            ),
        )
        return result

    def get_test_run(self, test_run_id: str) -> dict[str, Any]:
        row = self.database.fetch_one("SELECT * FROM test_runs WHERE id = ?", (test_run_id,))
        if row is None:
            raise KeyError("金标准测试记录不存在")
        row["details"] = json_load(row.pop("details_json"), [])
        return row
