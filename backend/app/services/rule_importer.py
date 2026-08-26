from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..db import Database

EXPECTED_HEADERS = [
    "规则编号",
    "启用状态",
    "场景",
    "适用工序",
    "触发条件",
    "具体要求",
    "数值或阈值",
    "规则效力",
    "危险源",
    "风险等级",
    "适用岗位",
    "检查方式",
    "规范名称",
    "标准编号",
    "标准层级",
    "标准状态",
    "条款编号",
    "原文",
    "PDF页码",
    "正文页码",
]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def import_rules(database: Database, workbook_path: Path) -> tuple[int, int]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"规则工作簿不存在：{workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "规则库" not in workbook.sheetnames:
        raise ValueError("规则工作簿缺少“规则库”工作表")
    worksheet = workbook["规则库"]

    header_row = None
    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row and _text(row[0]) == "规则编号":
            header_row = row_index
            actual_headers = [_text(value) for value in row[: len(EXPECTED_HEADERS)]]
            if actual_headers != EXPECTED_HEADERS:
                raise ValueError("规则库字段与首版数据契约不一致")
            break
    if header_row is None:
        raise ValueError("规则库中未找到字段标题行")

    now = datetime.now(timezone.utc).isoformat()
    rows: list[tuple[Any, ...]] = []
    skipped = 0
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        values = list(row[: len(EXPECTED_HEADERS)])
        rule_id = _text(values[0])
        if not rule_id:
            skipped += 1
            continue
        rows.append(
            (
                rule_id,
                *[_text(value) for value in values[1:18]],
                _integer(values[18]),
                _integer(values[19]),
                str(workbook_path),
                row_index,
                now,
            )
        )

    sql = """
    INSERT INTO audit_rules (
        rule_id, enabled_status, scene, process, trigger_condition, requirement, threshold,
        rule_effect, hazards, risk_level, roles, inspection_method, standard_name,
        standard_code, standard_level, standard_status, clause, original_text,
        pdf_page, printed_page, source_workbook, source_row, imported_at
    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    ON CONFLICT(rule_id) DO UPDATE SET
        enabled_status=excluded.enabled_status,
        scene=excluded.scene,
        process=excluded.process,
        trigger_condition=excluded.trigger_condition,
        requirement=excluded.requirement,
        threshold=excluded.threshold,
        rule_effect=excluded.rule_effect,
        hazards=excluded.hazards,
        risk_level=excluded.risk_level,
        roles=excluded.roles,
        inspection_method=excluded.inspection_method,
        standard_name=excluded.standard_name,
        standard_code=excluded.standard_code,
        standard_level=excluded.standard_level,
        standard_status=excluded.standard_status,
        clause=excluded.clause,
        original_text=excluded.original_text,
        pdf_page=excluded.pdf_page,
        printed_page=excluded.printed_page,
        source_workbook=excluded.source_workbook,
        source_row=excluded.source_row,
        imported_at=excluded.imported_at
    """
    database.executemany(sql, rows)
    workbook.close()
    return len(rows), skipped
