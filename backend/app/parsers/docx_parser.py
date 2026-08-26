from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class ParsedSegment:
    sequence_no: int
    segment_type: str
    heading_path: str
    location: str
    text: str
    markdown: str


@dataclass(frozen=True)
class ParsedDocument:
    segments: list[ParsedSegment]
    warnings: list[str]
    markdown: str


CHAPTER_HEADING_RE = re.compile(
    r"^第[一二三四五六七八九十百千万零〇两0-9]+(?P<unit>[篇章节])"
)
DECIMAL_HEADING_RE = re.compile(
    r"^(?P<number>\d+(?:\.\d+){1,4})(?:[、.．\s]|(?=[\u4e00-\u9fff]))"
)
SIMPLE_NUMBER_HEADING_RE = re.compile(r"^\d+[、.．]\s*\S+")
CHINESE_NUMBER_HEADING_RE = re.compile(
    r"^[一二三四五六七八九十百千万零〇两]+[、.．]\s*\S+"
)
SENTENCE_BOUNDARY_RE = re.compile(r"(?<=[。！？；])")
MEASURE_SENTENCE_RE = re.compile(
    r"应|必须|不得|严禁|不准|要求|停止|检查|采取|方可|才准|需|要|"
    r"戴好|系好|穿好|设置|配备|控制|禁止|做到|确保|防止"
)


def validate_docx(path: Path) -> None:
    if path.suffix.lower() != ".docx":
        raise ValueError("首版只支持 DOCX 施工方案")
    if not zipfile.is_zipfile(path):
        raise ValueError("文件扩展名为 DOCX，但内容不是有效的 Office Open XML 文档")
    try:
        Document(path)
    except Exception as exc:  # python-docx 的异常类型不稳定，统一转换业务错误。
        raise ValueError("DOCX 文件无法解析，可能已损坏或受密码保护") from exc


def _clean_text(text: str) -> str:
    return re.sub(r"[\t\u00a0 ]+", " ", text.replace("\r", "\n")).strip()


def _escape_markdown_cell(text: str) -> str:
    """Preserve table cell boundaries and line breaks in Markdown tables."""
    return _clean_text(text).replace("|", "\\|").replace("\n", "<br>")


def _markdown_table(rows: list[list[str]], table_index: int) -> list[str]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    lines = [f"<!-- source:T{table_index:03d} -->"]
    lines.append("| " + " | ".join(normalized[0]) + " |")
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return lines


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style else ""
    text = _clean_text(paragraph.text)
    if not text:
        return None

    # Numbering is more reliable than a frequently misused Word heading level. For
    # example, a paragraph styled Heading 1 but numbered “6.3” is still a level-2
    # section. Conversely, simple “1、……” body lists are not headings by default.
    chapter = CHAPTER_HEADING_RE.match(text)
    if chapter:
        return {"篇": 1, "章": 1, "节": 2}[chapter.group("unit")]
    decimal = DECIMAL_HEADING_RE.match(text)
    if decimal and len(text) <= 100:
        remainder = text[decimal.end() :].strip()
        # Legacy schemes often number every safety measure as 9.2.5/10.19 and
        # assign it a heading-like Word style. A complete normative sentence is
        # auditable body text, not an empty structural heading.
        if MEASURE_SENTENCE_RE.search(remainder):
            return None
        return min(decimal.group("number").count(".") + 1, 5)

    style_match = re.search(r"(?:Heading|标题)\s*([1-9])", style_name, re.IGNORECASE)
    if style_match and len(text) <= 100:
        return int(style_match.group(1))

    # Some legacy schemes have unstyled chapter names. Only promote short,
    # visually prominent numbered paragraphs; this prevents ordinary measure
    # lists from repeatedly resetting the chapter path.
    nonempty_runs = [run for run in paragraph.runs if _clean_text(run.text)]
    mostly_bold = bool(nonempty_runs) and sum(bool(run.bold) for run in nonempty_runs) >= (
        len(nonempty_runs) + 1
    ) // 2
    if len(text) <= 45 and mostly_bold:
        if SIMPLE_NUMBER_HEADING_RE.match(text):
            return 1
        if CHINESE_NUMBER_HEADING_RE.match(text):
            return 2
    return None


def _split_long_text(text: str, max_chars: int = 1200) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    sentences = [part for part in SENTENCE_BOUNDARY_RE.split(text) if part]
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        if current and len(current) + len(sentence) > max_chars:
            chunks.append(current)
            current = sentence
        else:
            current += sentence
    if current:
        chunks.append(current)
    return chunks or [text[index : index + max_chars] for index in range(0, len(text), max_chars)]


def parse_docx(path: Path) -> ParsedDocument:
    validate_docx(path)
    document = Document(path)
    headings: list[str] = []
    segments: list[ParsedSegment] = []
    warnings = ["DOCX 不提供可靠页码，首版使用标题路径和段落／表格坐标定位原文。"]
    paragraph_index = 0
    table_index = 0
    sequence_no = 0
    markdown_lines: list[str] = []

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = _clean_text(block.text)
            if not text:
                continue
            paragraph_index += 1
            level = _heading_level(block)
            if level is not None:
                headings = headings[: level - 1]
                while len(headings) < level - 1:
                    headings.append("未命名层级")
                headings.append(text)
                markdown_lines.extend([f"{'#' * level} {text}", ""])
                continue
            heading_path = " > ".join(headings) or "文档正文"
            parts = _split_long_text(text)
            markdown_lines.extend(
                [f"<!-- source:P{paragraph_index:03d} -->", text, ""]
            )
            for part_index, part in enumerate(parts, start=1):
                sequence_no += 1
                suffix = f"-{part_index}" if len(parts) > 1 else ""
                segments.append(
                    ParsedSegment(
                        sequence_no=sequence_no,
                        segment_type="paragraph",
                        heading_path=heading_path,
                        location=f"{heading_path}｜P{paragraph_index:03d}{suffix}",
                        text=part,
                        markdown=part,
                    )
                )
        elif isinstance(block, Table):
            table_index += 1
            heading_path = " > ".join(headings) or "文档正文"
            markdown_rows: list[list[str]] = []
            for row_index, row in enumerate(block.rows, start=1):
                cells = [_clean_text(cell.text) for cell in row.cells]
                markdown_cells = [_escape_markdown_cell(cell.text) for cell in row.cells]
                markdown_rows.append(markdown_cells)
                text = "｜".join(value for value in cells if value)
                if not text:
                    continue
                sequence_no += 1
                segments.append(
                    ParsedSegment(
                        sequence_no=sequence_no,
                        segment_type="table_row",
                        heading_path=heading_path,
                        location=f"{heading_path}｜T{table_index:03d}-R{row_index:03d}",
                        text=text,
                        markdown="| " + " | ".join(markdown_cells) + " |",
                    )
                )
            table_markdown = _markdown_table(markdown_rows, table_index)
            if table_markdown:
                markdown_lines.extend([*table_markdown, ""])

    if not segments:
        raise ValueError("DOCX 中未解析到可审计的正文或表格内容")
    return ParsedDocument(
        segments=segments,
        warnings=warnings,
        markdown="\n".join(markdown_lines).strip() + "\n",
    )
