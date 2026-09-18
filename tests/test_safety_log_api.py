from datetime import date
from base64 import b64decode
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from backend.app.main import create_app
from backend.app.services.safety_log_docx import render_safety_log_docx


def _create_task(client: TestClient) -> dict:
    response = client.post(
        "/worker-assistant/messages",
        json={
            "message": "今天下午在12层拆除外墙模板",
            "worker_ref": "项目作业人员",
            "team_ref": "模板班组",
            "use_llm": False,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_safety_log_is_versioned_downloadable_and_uses_current_project_name(test_settings):
    app = create_app(test_settings)
    today = date.today().isoformat()
    with TestClient(app) as client:
        _create_task(client)

        created = client.post("/api/v1/safety-logs", json={"assessment_date": today})
        assert created.status_code == 201, created.text
        first = created.json()
        assert first["project_name"] == "XX综合医院扩建项目"
        assert first["assessment_date"] == today
        assert first["version"] == 1
        assert first["version_created"] is True
        assert first["content"]["summary"]["task_count"] == 1
        assert first["content"]["source_modules"] == [
            "施工方案审查",
            "班前风险分析",
            "现场隐患巡检",
            "隐患整改记录",
        ]
        assert first["content"]["timeline"]
        assert any(item["type"] == "pre_job_risk" for item in first["content"]["timeline"])
        assert "open_items" in first["content"]
        assert "web_verification" not in first["content"]

        duplicate = client.post("/api/v1/safety-logs", json={"assessment_date": today})
        assert duplicate.status_code == 200, duplicate.text
        assert duplicate.json()["id"] == first["id"]
        assert duplicate.json()["version_created"] is False

        download = client.get(f"/api/v1/safety-logs/{first['id']}/download")
        assert download.status_code == 200, download.text
        assert download.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        word = Document(BytesIO(download.content))
        text = "\n".join(paragraph.text for paragraph in word.paragraphs)
        assert "XX综合医院扩建项目" in text
        assert "施工安全全过程日志" in text
        assert "天气摘要：" not in text
        assert "安全学习与培训建议" not in text
        assert "二、当日安全过程时间线" in text
        assert "三、施工方案审查与整改" in text
        assert "四、班前风险分析" in text
        assert "五、现场隐患巡检与整改" in text
        assert "六、截至当日未闭环事项" in text
        assert "七、数据来源与追溯说明" not in text
        assert "这是该日期的首个安全日志版本" not in text

        catalog = client.get("/api/v1/projects").json()
        project_id = catalog["active_project_id"]
        renamed = client.patch(
            f"/api/v1/projects/{project_id}", json={"name": "武汉市XX区住宅楼施工项目"}
        )
        assert renamed.status_code == 200, renamed.text

        renamed_log = client.post("/api/v1/safety-logs", json={"assessment_date": today})
        assert renamed_log.status_code == 201, renamed_log.text
        second = renamed_log.json()
        assert second["project_name"] == "武汉市XX区住宅楼施工项目"
        assert second["version"] == 2
        assert second["change"]["changed_sources"] == ["项目名称"]


def test_safety_log_docx_places_basis_and_before_after_images(tmp_path):
    image_bytes = b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    upload_root = tmp_path / "uploads"
    before_dir = upload_root / "hazard_images"
    after_dir = upload_root / "rectification_images"
    output_dir = upload_root / "safety_logs"
    before_dir.mkdir(parents=True)
    after_dir.mkdir(parents=True)
    (before_dir / "before.png").write_bytes(image_bytes)
    (after_dir / "after.png").write_bytes(image_bytes)
    target = output_dir / "safety-log.docx"
    log = {
        "project_name": "测试项目",
        "assessment_date": date.today().isoformat(),
        "version": 1,
        "created_at": "2026-09-17T08:00:00",
        "change": {},
        "content": {
            "summary": {
                "plan_activity_count": 0,
                "task_count": 0,
                "inspection_count": 1,
                "safety_item_count": 1,
                "closed_safety_item_count": 1,
                "open_item_count": 0,
            },
            "timeline": [],
            "plan_audits": [],
            "plan_revisions": [],
            "weather": {},
            "tasks": [],
            "onsite_inspections": [
                {
                    "id": "inspection-1",
                    "original_filename": "整改前.png",
                    "stored_image_name": "before.png",
                }
            ],
            "onsite_rectification": [
                {
                    "inspection_id": "inspection-1",
                    "item_no": "AQ-001",
                    "title": "挡脚板变形松脱",
                    "fact_description": "内侧脚手架挡脚板变形并松脱。",
                    "rectification_requirement": "更换合格挡脚板并固定。",
                    "basis": [
                        {
                            "standard_code": "GB 51210-2016",
                            "standard_name": "建筑施工脚手架安全技术统一标准",
                            "clause": "8.2.3",
                            "quote": "作业层应设置挡脚板。",
                        }
                    ],
                    "before_original_filename": "整改前.png",
                    "before_stored_image_name": "before.png",
                    "after_original_filename": "整改后.png",
                    "after_stored_image_name": "after.png",
                    "rectification_description": "已更换挡脚板。",
                    "status": "closed",
                    "order_status": "closed",
                    "latest_review_result": "pass",
                    "latest_review_reason": "复核通过",
                }
            ],
            "open_items": [],
        },
    }
    render_safety_log_docx(log, target)
    word = Document(target)
    all_text = "\n".join(
        [paragraph.text for paragraph in word.paragraphs]
        + [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    )
    assert "相关规范依据" in all_text
    assert "GB 51210-2016" in all_text
    assert "8.2.3" in all_text
    assert "七、数据来源与追溯说明" not in all_text
    assert len(word.inline_shapes) == 2


def test_safety_log_basis_text_is_concise():
    from backend.app.services.safety_log_docx import _basis_text

    text = _basis_text(
        [
            {
                "standard_code": "GB 51210-2016",
                "standard_name": "建筑施工脚手架安全技术统一标准",
                "clause": "3.2",
                "quote": "3.2 安全等级和安全系数\n"
                + "这是一段与隐患不直接相关的超长条文说明内容"
                * 8,
            }
        ]
    )
    assert text.startswith(
        "GB 51210-2016《建筑施工脚手架安全技术统一标准》 3.2"
    )
    assert len(text.splitlines()[-1]) <= 82


def test_safety_log_lists_history_and_agent_can_generate_it(test_settings):
    app = create_app(test_settings)
    today = date.today().isoformat()
    with TestClient(app) as client:
        _create_task(client)
        response = client.post(
            "/api/v1/agent/messages", json={"message": "生成今天的高处作业安全日志"}
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["agent_name"] == "safety_log_agent"
        assert body["metadata"]["tools"] == ["safety_log.generate"]

        history = client.get("/api/v1/safety-logs")
        assert history.status_code == 200
        assert history.json()[0]["assessment_date"] == today
        assert client.get("/api/v1/safety-logs/latest").json()["id"] == history.json()[0]["id"]
