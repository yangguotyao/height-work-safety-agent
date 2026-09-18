from docx import Document

from backend.app.parsers.docx_parser import parse_docx


def test_parses_docx_into_traceable_segments(sample_docx):
    parsed = parse_docx(sample_docx)

    assert len(parsed.segments) > 100
    assert parsed.warnings
    assert all(segment.location for segment in parsed.segments)
    assert any(segment.segment_type == "table_row" for segment in parsed.segments)
    assert any("脚手架" in segment.text for segment in parsed.segments)
    assert parsed.markdown
    assert "#" in parsed.markdown
    assert "|" in parsed.markdown
    assert "<!-- source:" in parsed.markdown


def test_does_not_treat_plain_numbered_measures_as_headings(tmp_path):
    path = tmp_path / "numbered-list.docx"
    document = Document()
    document.add_heading("第六章 脚手架施工", level=1)
    document.add_paragraph("1、作业人员应正确佩戴安全带。")
    document.add_paragraph("本段仍属于脚手架施工章节。")
    document.save(path)

    parsed = parse_docx(path)

    assert [segment.heading_path for segment in parsed.segments] == [
        "第六章 脚手架施工",
        "第六章 脚手架施工",
    ]


def test_numbering_corrects_incorrect_word_heading_level(tmp_path):
    path = tmp_path / "wrong-style-level.docx"
    document = Document()
    document.add_heading("第六章 脚手架施工", level=1)
    document.add_heading("6.3脚手架拆除", level=1)
    document.add_paragraph("拆除作业应按顺序实施。")
    document.save(path)

    parsed = parse_docx(path)

    assert parsed.segments[0].heading_path == "第六章 脚手架施工 > 6.3脚手架拆除"


def test_decimal_numbered_safety_measure_remains_searchable_body_text(tmp_path):
    path = tmp_path / "numbered-measure.docx"
    document = Document()
    document.add_heading("10、文明施工要求", level=1)
    document.add_heading(
        "10.19六级以上大风、大雪、大雾、大雨天气停止脚手架作业，"
        "雨雪后应检查并清扫脚手板。",
        level=2,
    )
    document.save(path)

    parsed = parse_docx(path)

    assert len(parsed.segments) == 1
    assert "雨雪后应检查" in parsed.segments[0].text
    assert parsed.segments[0].heading_path == "10、文明施工要求"


def test_markdown_table_keeps_empty_cells_and_line_breaks(tmp_path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).text = "栏杆"
    table.cell(1, 1).text = "上杆1.2m\n下杆0.6m"
    table.cell(1, 2).text = ""
    document.save(path)

    parsed = parse_docx(path)

    assert "| 栏杆 | 上杆1.2m<br>下杆0.6m |  |" in parsed.markdown
    assert parsed.segments[1].markdown.endswith("|  |")


def test_recognizes_legacy_scaffold_calculation_boundaries(tmp_path):
    path = tmp_path / "legacy-headings.docx"
    document = Document()
    document.add_heading("第一节 外脚手架搭设要求", level=2)
    document.add_paragraph("外脚手架采用双排架。")
    document.add_paragraph("内脚手架及砖架搭设")
    document.add_paragraph("内脚手架采用满堂架。")
    document.add_paragraph("脚手架计算书")
    document.add_paragraph("悬挑式钢管脚手架计算书")
    document.add_paragraph("立杆纵距1.5m。")
    document.save(path)

    parsed = parse_docx(path)

    assert parsed.segments[1].heading_path.endswith("内脚手架及砖架搭设")
    assert parsed.segments[2].heading_path == "脚手架计算书 > 悬挑式钢管脚手架计算书"
